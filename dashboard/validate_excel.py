"""Compare the Python engine with the cached values of the sample Excel's DYNMIC sheet (year in cell BE2).

    python validate_excel.py

1) Engine check: tests/mashhad_excel_check.yaml feeds the engine the same inputs the Excel used
   -> every cell should match (only the aquifer areas differ slightly: shapefile vs boundary_area sheet).
2) Data-source check: configs/mashhad.yaml (Amar_SAD + CHAH + Simaye Ab PDFs + shapefiles)
   -> shows which rows/columns differ from the sample Excel and by how much.
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from wa.build import get_base  # noqa: E402
from wa.config import load_config, resolve  # noqa: E402
from wa.engine import compute  # noqa: E402
from wa.table import CTRL_COL, G_COL, KR, KW, ROW_BAL, ROW_CONS, ROW_TOTAL, SUMMARY, build_table  # noqa: E402
from wa.text import norm  # noqa: E402

pd.set_option("display.width", 220)


def compare(cfg, ws):
    year = norm(ws["BE2"].value)
    base = get_base(cfg, log=lambda *_: None)
    uses = cfg["uses"]
    W, RET, RES = compute(base, cfg, uses, "attribute")
    counties = sorted(set(W.county) | set(RES.county))
    t, totals = build_table(base, cfg, uses, W, RET, RES, year, base["meta"]["sub_basins"], counties,
                            [u["name"] for u in uses], "excel", ["dam", "river", "gw", "transfer", "treatment"], True)
    sb_label = cfg["spatial"]["sub_basin_label"]

    def row_key(r):
        f = ws.cell(r, 6).value
        if r == 21:
            return ("تصفیه خانه", cfg["treatment_label"])
        if isinstance(f, (int, float)):
            return ("آب زیرزمینی", cfg["spatial"]["aquifer_label"].format(id=int(f)))
        f = norm(f)
        if f.startswith("بازه رودخانه"):
            return ("بازه رودخانه", f)
        return ("انتقالی", f) if r == 20 else ("سد", f)

    def cell(key, col):
        return float(t.loc[key, col]) if key in t.index and col in t.columns and pd.notna(t.loc[key, col]) else 0.0

    hdr = {u["name"]: next(c[0] for c in t.columns if c[0].startswith(u["name"] + " ")) for u in uses}
    rows, cur = [], None
    for c in range(8, 54, 2):
        if ws.cell(2, c).value:
            cur = norm(ws.cell(2, c).value)
        m = ws.cell(3, c).value
        member = "کل" if m is None else (sb_label.format(id=int(m)) if isinstance(m, (int, float)) else norm(m))
        h = hdr.get(cur)
        for r in list(range(7, 22)) + [22, 23]:
            key = (SUMMARY, ROW_TOTAL if r == 22 else ROW_CONS) if r >= 22 else row_key(r)
            for kind, col in ((KW, c), (KR, c + 1)):
                xv = ws.cell(r, col).value
                if r >= 22 and xv is None:
                    continue
                pv = cell(key, (h, member, kind)) if h else 0.0
                rows.append((ws.cell(r, col).coordinate, key[1], cur, member, kind, float(xv or 0), pv))
    for r in range(7, 21):
        key = row_key(r)
        for col, tc in ((7, G_COL), (54, CTRL_COL)):
            rows.append((ws.cell(r, col).coordinate, key[1], "-", "-", tc[0], float(ws.cell(r, col).value or 0),
                         cell(key, tc)))
    rows.append(("H24", ROW_BAL, "-", "-", "-", float(ws["H24"].value), totals["balance"]))
    df = pd.DataFrame(rows, columns=["cell", "row", "use", "member", "kind", "excel", "python"])
    df["diff"] = df.python - df.excel
    return year, df


ws = openpyxl.load_workbook(resolve(load_config(HERE / "configs" / "mashhad.yaml"), "water_acounting_1Sep25.xlsx"),
                            data_only=True)["DYNMIC"]

print("=" * 90, "\n1) آزمون موتور محاسبات با همان ورودی‌های اکسل نمونه")
year, df = compare(load_config(HERE / "tests" / "mashhad_excel_check.yaml"), ws)
bad = df[df["diff"].abs() > 0.01]
print(f"سال آبی {year} | سلول‌ها: {len(df)} | اختلاف > 0.01: {len(bad)} | بیشترین اختلاف: {df['diff'].abs().max():.4f}")
if len(bad):
    print(bad.to_string(index=False))

print("=" * 90, "\n2) داشبورد اصلی (Amar_SAD + CHAH_MASHAHD + سیمای آب + شیپ‌فایل) در مقایسه با اکسل نمونه")
year, df = compare(load_config(HERE / "configs" / "mashhad.yaml"), ws)
bad = df[df["diff"].abs() > 0.01]
print(f"سلول‌ها: {len(df)} | اختلاف > 0.01: {len(bad)}")
if len(bad):
    by_row = bad.groupby("row")["diff"].agg(["count", "sum"]).round(3).sort_values("sum", key=abs, ascending=False)
    print(by_row.to_string())
    print(bad.reindex(bad["diff"].abs().sort_values(ascending=False).index).head(15).to_string(index=False))
